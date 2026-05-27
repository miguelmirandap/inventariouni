from __future__ import annotations

import csv
import io
import os
import shutil
import sqlite3
import sys
import threading
from datetime import datetime
from pathlib import Path
from functools import wraps

from flask import Flask, Response, flash, redirect, render_template, request, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash

from app_inventario import init_db, seed_from_csv_if_empty

if getattr(sys, "frozen", False):
    APP_DIR = Path(sys.executable).resolve().parent
    RESOURCE_DIR = Path(getattr(sys, "_MEIPASS", APP_DIR))
else:
    APP_DIR = Path(__file__).resolve().parent
    RESOURCE_DIR = APP_DIR

DB_PATH = APP_DIR / "inventario.db"
BACKUP_DIR = APP_DIR / "backups"

app = Flask(
    __name__,
    template_folder=str(RESOURCE_DIR / "templates"),
    static_folder=str(RESOURCE_DIR / "static"),
)
app.config["SECRET_KEY"] = os.environ.get(
    "SECRET_KEY", "inventario-uni-2026-cambiar-en-produccion"
)


def format_currency(value):
    try:
        amount = float(value)
    except (TypeError, ValueError):
        return "$0,00"

    negative = amount < 0
    amount = abs(amount)
    formatted = f"{amount:,.2f}"
    formatted = formatted.replace(",", "X").replace(".", ",").replace("X", ".")
    if negative:
        formatted = f"-{formatted}"
    return f"${formatted}"


@app.template_filter("currency")
def currency_filter(value):
    return format_currency(value)


def parse_float(value: str, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def parse_int(value: str, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def current_date() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def is_password_hashed(value: str) -> bool:
    return value.startswith("scrypt:") or value.startswith("pbkdf2:")


def connect_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _migrate_users_table_if_needed(conn: sqlite3.Connection) -> None:
    # Existing installations may have role check limited to admin/caja.
    create_sql_row = conn.execute(
        "SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'usuarios'"
    ).fetchone()
    if not create_sql_row or not create_sql_row[0]:
        return

    create_sql = create_sql_row[0]
    if "supervisor" in create_sql:
        return

    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS usuarios_new (
            username TEXT PRIMARY KEY,
            password TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('admin', 'caja', 'supervisor')),
            nombre TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        INSERT OR REPLACE INTO usuarios_new (username, password, role, nombre)
        SELECT username, password, role, nombre FROM usuarios
        """
    )
    conn.execute("DROP TABLE usuarios")
    conn.execute("ALTER TABLE usuarios_new RENAME TO usuarios")
    conn.commit()


def ensure_users_table() -> None:
    conn = connect_db()
    _migrate_users_table_if_needed(conn)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS usuarios (
            username TEXT PRIMARY KEY,
            password TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('admin', 'caja', 'supervisor')),
            nombre TEXT NOT NULL
        )
        """
    )

    conn.execute(
        """
        INSERT OR IGNORE INTO usuarios (username, password, role, nombre)
        VALUES (?, ?, 'admin', 'Administrador')
        """,
        ("admin", generate_password_hash("admin123")),
    )
    conn.execute(
        """
        INSERT OR IGNORE INTO usuarios (username, password, role, nombre)
        VALUES (?, ?, 'caja', 'Caja')
        """,
        ("caja", generate_password_hash("caja123")),
    )
    conn.execute(
        """
        INSERT OR IGNORE INTO usuarios (username, password, role, nombre)
        VALUES (?, ?, 'supervisor', 'Supervisor')
        """,
        ("supervisor", generate_password_hash("super123")),
    )

    users = conn.execute("SELECT username, password FROM usuarios").fetchall()
    for user in users:
        if not is_password_hashed(user["password"]):
            conn.execute(
                "UPDATE usuarios SET password = ? WHERE username = ?",
                (generate_password_hash(user["password"]), user["username"]),
            )

    conn.commit()
    conn.close()


def ensure_backup_dir() -> None:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)


def backup_filename_for_today() -> Path:
    return BACKUP_DIR / f"inventario_{current_date().replace('-', '')}.db"


def create_backup() -> Path:
    ensure_backup_dir()
    target = backup_filename_for_today()
    shutil.copy2(DB_PATH, target)
    return target


def ensure_daily_backup() -> None:
    ensure_backup_dir()
    target = backup_filename_for_today()
    if not target.exists() and DB_PATH.exists():
        shutil.copy2(DB_PATH, target)


def ensure_cartera_tables() -> None:
    conn = connect_db()
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS clientes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            telefono TEXT,
            correo TEXT,
            notas TEXT DEFAULT '',
            activo INTEGER NOT NULL DEFAULT 1
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS deudas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cliente_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            fecha_vencimiento TEXT,
            concepto TEXT NOT NULL,
            monto_total REAL NOT NULL,
            monto_pagado REAL NOT NULL DEFAULT 0,
            estado TEXT NOT NULL DEFAULT 'PENDIENTE' CHECK(estado IN ('PENDIENTE', 'PAGADA', 'VENCIDA')),
            usuario TEXT,
            FOREIGN KEY (cliente_id) REFERENCES clientes(id)
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS abonos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            deuda_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            monto REAL NOT NULL,
            metodo_pago TEXT,
            usuario TEXT,
            FOREIGN KEY (deuda_id) REFERENCES deudas(id)
        )
        """
    )
    conn.commit()
    conn.close()


def ensure_ventas_tables() -> None:
    conn = connect_db()
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS ventas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            tipo_pago TEXT NOT NULL CHECK(tipo_pago IN ('CONTADO', 'FIADO')),
            cliente_id INTEGER,
            subtotal REAL NOT NULL,
            total REAL NOT NULL,
            notas TEXT DEFAULT '',
            usuario TEXT,
            FOREIGN KEY (cliente_id) REFERENCES clientes(id)
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS venta_detalle (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            venta_id INTEGER NOT NULL,
            sku TEXT NOT NULL,
            producto TEXT NOT NULL,
            cantidad INTEGER NOT NULL,
            precio_unitario REAL NOT NULL,
            total_linea REAL NOT NULL,
            FOREIGN KEY (venta_id) REFERENCES ventas(id),
            FOREIGN KEY (sku) REFERENCES productos(sku)
        )
        """
    )
    conn.commit()
    conn.close()


def registrar_salida_por_venta(
    conn: sqlite3.Connection,
    producto: sqlite3.Row,
    cantidad: int,
    precio_unitario: float,
    responsable: str,
    motivo: str,
) -> None:
    nuevo_stock = int(producto["stock_actual"]) - cantidad
    if nuevo_stock < 0:
        raise ValueError("Stock insuficiente para completar la venta.")

    fecha = current_date()
    costo_unitario = float(producto["precio_compra"])
    conn.execute(
        """
        INSERT INTO movimientos (
            fecha, tipo, sku, producto, cantidad, costo_unitario,
            precio_unitario, motivo, responsable
        ) VALUES (?, 'SALIDA', ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            fecha,
            producto["sku"],
            producto["nombre"],
            -cantidad,
            costo_unitario,
            precio_unitario,
            motivo,
            responsable,
        ),
    )

    conn.execute(
        """
        UPDATE productos
        SET stock_actual = ?,
            precio_venta = ?,
            fecha_ultima_salida = ?
        WHERE sku = ?
        """,
        (nuevo_stock, precio_unitario, fecha, producto["sku"]),
    )


def bootstrap() -> None:
    conn = connect_db()
    init_db(conn)
    seed_from_csv_if_empty(conn)
    conn.close()
    ensure_users_table()
    ensure_cartera_tables()
    ensure_ventas_tables()
    ensure_daily_backup()


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "username" not in session:
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped


def role_required(*roles):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if session.get("role") not in roles:
                flash("No tienes permisos para esta accion.", "error")
                return redirect(url_for("dashboard"))
            return view(*args, **kwargs)

        return wrapped

    return decorator


def has_any_role(*roles: str) -> bool:
    return session.get("role") in set(roles)


@app.route("/")
def index():
    if "username" in session:
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()

        conn = connect_db()
        user = conn.execute(
            "SELECT * FROM usuarios WHERE username = ?",
            (username,),
        ).fetchone()
        conn.close()

        if not user or not check_password_hash(user["password"], password):
            flash("Usuario o contrasena incorrectos.", "error")
            response = app.make_response(render_template("login.html"))
            response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
            response.headers["Pragma"] = "no-cache"
            response.headers["Expires"] = "0"
            return response

        session["username"] = user["username"]
        session["role"] = user["role"]
        session["nombre"] = user["nombre"]
        flash(f"Bienvenido, {user['nombre']}.", "ok")
        return redirect(url_for("dashboard"))

    response = app.make_response(render_template("login.html"))
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.route("/perfil/password", methods=["GET", "POST"])
@login_required
def perfil_password():
    if request.method == "POST":
        actual = request.form.get("password_actual", "").strip()
        nueva = request.form.get("password_nueva", "").strip()
        confirmar = request.form.get("password_confirmar", "").strip()

        if len(nueva) < 6:
            flash("La nueva contrasena debe tener al menos 6 caracteres.", "error")
            return redirect(url_for("perfil_password"))
        if nueva != confirmar:
            flash("La confirmacion de contrasena no coincide.", "error")
            return redirect(url_for("perfil_password"))

        conn = connect_db()
        user = conn.execute(
            "SELECT * FROM usuarios WHERE username = ?",
            (session.get("username", ""),),
        ).fetchone()
        if not user or not check_password_hash(user["password"], actual):
            conn.close()
            flash("Contrasena actual incorrecta.", "error")
            return redirect(url_for("perfil_password"))

        conn.execute(
            "UPDATE usuarios SET password = ? WHERE username = ?",
            (generate_password_hash(nueva), session.get("username", "")),
        )
        conn.commit()
        conn.close()
        flash("Contrasena actualizada correctamente.", "ok")
        return redirect(url_for("dashboard"))

    return render_template("perfil_password.html")


@app.route("/logout")
@login_required
def logout():
    session.clear()
    flash("Sesion cerrada.", "ok")
    response = redirect(url_for("login"))
    response.set_cookie("session", "", expires=0)
    return response


@app.route("/dashboard")
@login_required
def dashboard():
    conn = connect_db()
    resumen = conn.execute(
        """
        SELECT
            COUNT(*) AS total_productos,
            COALESCE(SUM(stock_actual * precio_compra), 0) AS valor_inventario,
            COALESCE(SUM(stock_actual * (precio_venta - precio_compra)), 0) AS utilidad_potencial
        FROM productos
        """
    ).fetchone()

    stock_bajo = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM productos
        WHERE stock_actual <= stock_minimo
        """
    ).fetchone()["total"]

    cartera_pendiente = conn.execute(
        """
        SELECT COALESCE(SUM(monto_total - monto_pagado), 0) AS total
        FROM deudas
        WHERE estado <> 'PAGADA'
        """
    ).fetchone()["total"]

    stock_critico = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM productos
        WHERE stock_actual <= CASE
            WHEN stock_minimo > 0 THEN MAX(1, CAST(stock_minimo * 0.5 AS INTEGER))
            ELSE 0
        END
        """
    ).fetchone()["total"]

    deudas_vencidas = conn.execute(
        """
        SELECT COUNT(*) AS total
        FROM deudas
        WHERE estado = 'VENCIDA' OR (
            estado = 'PENDIENTE' AND
            fecha_vencimiento IS NOT NULL AND
            fecha_vencimiento <> '' AND
            date(fecha_vencimiento) < date('now')
        )
        """
    ).fetchone()["total"]

    alertas = []
    if stock_critico > 0:
        alertas.append(f"Stock critico en {stock_critico} producto(s).")
    if deudas_vencidas > 0:
        alertas.append(f"Hay {deudas_vencidas} deuda(s) vencida(s).")

    conn.close()

    return render_template(
        "dashboard.html",
        resumen=resumen,
        stock_bajo=stock_bajo,
        cartera_pendiente=cartera_pendiente,
        stock_critico=stock_critico,
        deudas_vencidas=deudas_vencidas,
        alertas=alertas,
    )


@app.route("/alertas")
@login_required
def alertas():
    conn = connect_db()
    stock_rows = conn.execute(
        """
        SELECT sku, nombre, stock_actual, stock_minimo
        FROM productos
        WHERE stock_actual <= CASE
            WHEN stock_minimo > 0 THEN MAX(1, CAST(stock_minimo * 0.5 AS INTEGER))
            ELSE 0
        END
        ORDER BY stock_actual ASC
        """
    ).fetchall()

    deuda_rows = conn.execute(
        """
        SELECT
            d.id,
            c.nombre AS cliente,
            d.concepto,
            d.fecha_vencimiento,
            (d.monto_total - d.monto_pagado) AS saldo,
            CAST(julianday('now') - julianday(d.fecha_vencimiento) AS INTEGER) AS dias_atraso
        FROM deudas d
        JOIN clientes c ON c.id = d.cliente_id
        WHERE d.fecha_vencimiento IS NOT NULL
          AND d.fecha_vencimiento <> ''
          AND date(d.fecha_vencimiento) < date('now')
          AND (d.monto_total - d.monto_pagado) > 0
        ORDER BY dias_atraso DESC, saldo DESC
        """
    ).fetchall()
    conn.close()

    return render_template("alertas.html", stock_rows=stock_rows, deuda_rows=deuda_rows)


@app.route("/productos")
@login_required
def productos():
    q = request.args.get("q", "").strip()
    conn = connect_db()

    if q:
        rows = conn.execute(
            """
            SELECT *
            FROM productos
            WHERE sku LIKE ? OR nombre LIKE ? OR categoria LIKE ?
            ORDER BY nombre
            """,
            (f"%{q}%", f"%{q}%", f"%{q}%"),
        ).fetchall()
    else:
        rows = conn.execute("SELECT * FROM productos ORDER BY nombre").fetchall()

    conn.close()
    return render_template("productos.html", productos=rows, q=q)


@app.route("/productos/nuevo", methods=["POST"])
@login_required
@role_required("admin")
def nuevo_producto():
    data = request.form
    sku = data.get("sku", "").strip().upper()
    if not sku:
        flash("El SKU es obligatorio.", "error")
        return redirect(url_for("productos"))

    conn = connect_db()
    exists = conn.execute("SELECT 1 FROM productos WHERE sku = ?", (sku,)).fetchone()
    if exists:
        conn.close()
        flash("El SKU ya existe.", "error")
        return redirect(url_for("productos"))

    conn.execute(
        """
        INSERT INTO productos (
            sku, nombre, categoria, unidad, stock_actual, stock_minimo,
            precio_compra, precio_venta, proveedor, ubicacion,
            fecha_ultima_entrada, fecha_ultima_salida, observaciones
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '', '', ?)
        """,
        (
            sku,
            data.get("nombre", "").strip(),
            data.get("categoria", "General").strip() or "General",
            data.get("unidad", "pza").strip() or "pza",
            int(data.get("stock_actual", "0") or 0),
            int(data.get("stock_minimo", "0") or 0),
            float(data.get("precio_compra", "0") or 0),
            float(data.get("precio_venta", "0") or 0),
            data.get("proveedor", "INTERNO").strip() or "INTERNO",
            data.get("ubicacion", "Sin asignar").strip() or "Sin asignar",
            data.get("observaciones", "").strip(),
        ),
    )
    conn.commit()
    conn.close()
    flash("Producto creado correctamente.", "ok")
    return redirect(url_for("productos"))


@app.route("/movimientos", methods=["GET", "POST"])
@login_required
def movimientos():
    conn = connect_db()

    if request.method == "POST":
        tipo = request.form.get("tipo", "").strip().upper()
        sku = request.form.get("sku", "").strip().upper()
        cantidad = int(request.form.get("cantidad", "0") or 0)
        costo = float(request.form.get("costo_unitario", "0") or 0)
        precio = float(request.form.get("precio_unitario", "0") or 0)
        motivo = request.form.get("motivo", "").strip()

        if tipo not in {"ENTRADA", "SALIDA", "AJUSTE"}:
            conn.close()
            flash("Tipo de movimiento invalido.", "error")
            return redirect(url_for("movimientos"))

        producto = conn.execute("SELECT * FROM productos WHERE sku = ?", (sku,)).fetchone()
        if not producto:
            conn.close()
            flash("SKU no encontrado.", "error")
            return redirect(url_for("movimientos"))

        if tipo in {"ENTRADA", "SALIDA"} and cantidad <= 0:
            conn.close()
            flash("La cantidad debe ser mayor a 0.", "error")
            return redirect(url_for("movimientos"))

        if tipo == "ENTRADA":
            delta = cantidad
        elif tipo == "SALIDA":
            delta = -cantidad
        else:
            delta = cantidad

        nuevo_stock = producto["stock_actual"] + delta
        if nuevo_stock < 0:
            conn.close()
            flash("El movimiento deja stock negativo.", "error")
            return redirect(url_for("movimientos"))

        fecha = datetime.now().strftime("%Y-%m-%d")
        responsable = session.get("username", "")
        costo_final = costo if costo > 0 else float(producto["precio_compra"])
        precio_final = precio if precio > 0 else float(producto["precio_venta"])

        conn.execute(
            """
            INSERT INTO movimientos (
                fecha, tipo, sku, producto, cantidad, costo_unitario,
                precio_unitario, motivo, responsable
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                fecha,
                tipo,
                sku,
                producto["nombre"],
                delta,
                costo_final,
                precio_final,
                motivo,
                responsable,
            ),
        )

        fecha_entrada = producto["fecha_ultima_entrada"]
        fecha_salida = producto["fecha_ultima_salida"]
        if delta > 0:
            fecha_entrada = fecha
        if delta < 0:
            fecha_salida = fecha

        conn.execute(
            """
            UPDATE productos
            SET stock_actual = ?,
                precio_compra = ?,
                precio_venta = ?,
                fecha_ultima_entrada = ?,
                fecha_ultima_salida = ?
            WHERE sku = ?
            """,
            (nuevo_stock, costo_final, precio_final, fecha_entrada, fecha_salida, sku),
        )
        conn.commit()
        flash("Movimiento registrado.", "ok")

    productos_rows = conn.execute(
        "SELECT sku, nombre, stock_actual, stock_minimo FROM productos ORDER BY nombre"
    ).fetchall()
    ultimos = conn.execute(
        """
        SELECT fecha, tipo, sku, producto, cantidad, costo_unitario, precio_unitario, responsable
        FROM movimientos
        ORDER BY id DESC
        LIMIT 30
        """
    ).fetchall()
    conn.close()

    return render_template("movimientos.html", productos=productos_rows, movimientos=ultimos)


@app.route("/reportes")
@login_required
def reportes():
    fecha_inicio = request.args.get("fecha_inicio", "").strip()
    fecha_fin = request.args.get("fecha_fin", "").strip()

    if not fecha_inicio:
        fecha_inicio = datetime.now().replace(day=1).strftime("%Y-%m-%d")
    if not fecha_fin:
        fecha_fin = datetime.now().strftime("%Y-%m-%d")

    conn = connect_db()

    por_categoria = conn.execute(
        """
        SELECT
            p.categoria AS categoria,
            COALESCE(SUM(CASE WHEN m.tipo = 'SALIDA' THEN -m.cantidad ELSE 0 END), 0) AS unidades_vendidas,
            COALESCE(SUM(CASE WHEN m.tipo = 'SALIDA' THEN (-m.cantidad * m.precio_unitario) ELSE 0 END), 0) AS ingresos
        FROM movimientos m
        JOIN productos p ON p.sku = m.sku
        WHERE m.fecha BETWEEN ? AND ?
        GROUP BY p.categoria
        ORDER BY ingresos DESC
        """,
        (fecha_inicio, fecha_fin),
    ).fetchall()

    detalle = conn.execute(
        """
        SELECT
            m.fecha,
            m.tipo,
            m.sku,
            m.producto,
            m.cantidad,
            m.precio_unitario,
            p.categoria,
            (m.cantidad * m.precio_unitario) AS total
        FROM movimientos m
        JOIN productos p ON p.sku = m.sku
        WHERE m.fecha BETWEEN ? AND ?
        ORDER BY m.fecha DESC, m.id DESC
        """,
        (fecha_inicio, fecha_fin),
    ).fetchall()

    conn.close()

    return render_template(
        "reportes.html",
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        por_categoria=por_categoria,
        detalle=detalle,
    )


@app.route("/ventas", methods=["GET", "POST"])
@login_required
@role_required("admin", "supervisor", "caja")
def ventas():
    conn = connect_db()

    if request.method == "POST":
        tipo_pago = request.form.get("tipo_pago", "CONTADO").strip().upper()
        cliente_id_raw = request.form.get("cliente_id", "").strip()
        notas = request.form.get("notas", "").strip()

        skus = request.form.getlist("sku[]")
        cantidades = request.form.getlist("cantidad[]")
        precios = request.form.getlist("precio_unitario[]")

        if not skus and request.form.get("sku"):
            # Backward compatibility with single-line form submissions.
            skus = [request.form.get("sku", "")]
            cantidades = [request.form.get("cantidad", "0")]
            precios = [request.form.get("precio_unitario", "0")]

        lineas_raw: list[dict[str, object]] = []
        max_len = max(len(skus), len(cantidades), len(precios)) if skus else 0
        for i in range(max_len):
            sku = (skus[i] if i < len(skus) else "").strip().upper()
            if not sku:
                continue
            cantidad = parse_int(cantidades[i] if i < len(cantidades) else "0", 0)
            precio = parse_float(precios[i] if i < len(precios) else "0", 0.0)
            lineas_raw.append({"sku": sku, "cantidad": cantidad, "precio": precio})

        if not lineas_raw:
            conn.close()
            flash("Agrega al menos un producto a la venta.", "error")
            return redirect(url_for("ventas"))

        for linea in lineas_raw:
            if int(linea["cantidad"]) <= 0:
                conn.close()
                flash("Todas las cantidades deben ser mayores a 0.", "error")
                return redirect(url_for("ventas"))

        if tipo_pago not in {"CONTADO", "FIADO"}:
            conn.close()
            flash("Tipo de pago invalido.", "error")
            return redirect(url_for("ventas"))

        cliente_id = None
        if tipo_pago == "FIADO":
            if not cliente_id_raw:
                conn.close()
                flash("Para venta fiada debes seleccionar un cliente.", "error")
                return redirect(url_for("ventas"))
            cliente_id = parse_int(cliente_id_raw, 0)
            if cliente_id <= 0:
                conn.close()
                flash("Cliente invalido.", "error")
                return redirect(url_for("ventas"))

        sku_unicos = sorted({str(linea["sku"]) for linea in lineas_raw})
        placeholders = ",".join("?" for _ in sku_unicos)
        productos_db = conn.execute(
            f"SELECT * FROM productos WHERE sku IN ({placeholders})",
            sku_unicos,
        ).fetchall()
        productos_map = {row["sku"]: row for row in productos_db}

        for sku in sku_unicos:
            if sku not in productos_map:
                conn.close()
                flash(f"Producto no encontrado: {sku}", "error")
                return redirect(url_for("ventas"))

        cantidades_por_sku: dict[str, int] = {}
        for linea in lineas_raw:
            sku = str(linea["sku"])
            cantidad = int(linea["cantidad"])
            cantidades_por_sku[sku] = cantidades_por_sku.get(sku, 0) + cantidad

        for sku, cantidad in cantidades_por_sku.items():
            stock = int(productos_map[sku]["stock_actual"])
            if stock < cantidad:
                conn.close()
                flash(f"Stock insuficiente para {sku}. Disponible: {stock}", "error")
                return redirect(url_for("ventas"))

        lineas_finales: list[dict[str, object]] = []
        subtotal = 0.0
        for linea in lineas_raw:
            sku = str(linea["sku"])
            cantidad = int(linea["cantidad"])
            producto = productos_map[sku]
            precio = float(linea["precio"])
            if precio <= 0:
                precio = float(producto["precio_venta"])
            if precio <= 0:
                conn.close()
                flash(f"Precio invalido para {sku}.", "error")
                return redirect(url_for("ventas"))

            total_linea = round(cantidad * precio, 2)
            subtotal += total_linea
            lineas_finales.append(
                {
                    "sku": sku,
                    "cantidad": cantidad,
                    "precio": precio,
                    "total_linea": total_linea,
                    "producto": producto,
                }
            )

        subtotal = round(subtotal, 2)
        total = subtotal
        fecha = current_date()
        usuario = session.get("username", "")
        venta_id: int | None = None

        try:
            conn.execute("BEGIN")
            cur = conn.execute(
                """
                INSERT INTO ventas (fecha, tipo_pago, cliente_id, subtotal, total, notas, usuario)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (fecha, tipo_pago, cliente_id, subtotal, total, notas, usuario),
            )
            venta_id = int(cur.lastrowid)

            for linea in lineas_finales:
                producto = linea["producto"]
                conn.execute(
                    """
                    INSERT INTO venta_detalle (
                        venta_id, sku, producto, cantidad, precio_unitario, total_linea
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        venta_id,
                        linea["sku"],
                        producto["nombre"],
                        linea["cantidad"],
                        linea["precio"],
                        linea["total_linea"],
                    ),
                )

                registrar_salida_por_venta(
                    conn,
                    producto,
                    int(linea["cantidad"]),
                    float(linea["precio"]),
                    usuario,
                    f"Venta #{venta_id}",
                )

            if tipo_pago == "FIADO" and cliente_id:
                conn.execute(
                    """
                    INSERT INTO deudas (
                        cliente_id, fecha, fecha_vencimiento, concepto, monto_total,
                        monto_pagado, estado, usuario
                    ) VALUES (?, ?, NULL, ?, ?, 0, 'PENDIENTE', ?)
                    """,
                    (
                        cliente_id,
                        fecha,
                        f"Venta #{venta_id} (fiado)",
                        total,
                        usuario,
                    ),
                )

            conn.commit()
            flash(f"Venta #{venta_id} registrada correctamente.", "ok")
            conn.close()
            return redirect(url_for("venta_ticket", venta_id=venta_id))
        except Exception:
            conn.rollback()
            flash("No se pudo registrar la venta.", "error")

    productos_rows = conn.execute(
        """
        SELECT sku, nombre, stock_actual, precio_venta
        FROM productos
        WHERE stock_actual > 0
        ORDER BY nombre
        """
    ).fetchall()
    clientes_rows = conn.execute(
        "SELECT id, nombre FROM clientes WHERE activo = 1 ORDER BY nombre"
    ).fetchall()
    ventas_rows = conn.execute(
        """
        SELECT
            v.id,
            v.fecha,
            v.tipo_pago,
            COALESCE(c.nombre, '-') AS cliente,
            v.total,
            v.usuario,
            COALESCE(SUM(vd.cantidad), 0) AS items
        FROM ventas v
        JOIN venta_detalle vd ON vd.venta_id = v.id
        LEFT JOIN clientes c ON c.id = v.cliente_id
        GROUP BY v.id, v.fecha, v.tipo_pago, c.nombre, v.total, v.usuario
        ORDER BY v.id DESC
        LIMIT 40
        """
    ).fetchall()

    total_hoy_row = conn.execute(
        """
        SELECT
            COALESCE(SUM(total), 0) AS total,
            COALESCE(SUM(CASE WHEN tipo_pago = 'CONTADO' THEN total ELSE 0 END), 0) AS contado,
            COALESCE(SUM(CASE WHEN tipo_pago = 'FIADO' THEN total ELSE 0 END), 0) AS fiado
        FROM ventas
        WHERE fecha = ?
        """,
        (current_date(),),
    ).fetchone()
    conn.close()

    return render_template(
        "ventas.html",
        productos=productos_rows,
        clientes=clientes_rows,
        ventas=ventas_rows,
        total_hoy=total_hoy_row["total"],
        contado_hoy=total_hoy_row["contado"],
        fiado_hoy=total_hoy_row["fiado"],
    )


@app.route("/ventas/ticket/<int:venta_id>")
@login_required
def venta_ticket(venta_id: int):
    conn = connect_db()
    venta = conn.execute(
        """
        SELECT
            v.id,
            v.fecha,
            v.tipo_pago,
            v.subtotal,
            v.total,
            v.notas,
            v.usuario,
            COALESCE(c.nombre, 'Publico general') AS cliente
        FROM ventas v
        LEFT JOIN clientes c ON c.id = v.cliente_id
        WHERE v.id = ?
        """,
        (venta_id,),
    ).fetchone()

    if not venta:
        conn.close()
        flash("No existe ese ticket de venta.", "error")
        return redirect(url_for("ventas"))

    detalles = conn.execute(
        """
        SELECT sku, producto, cantidad, precio_unitario, total_linea
        FROM venta_detalle
        WHERE venta_id = ?
        ORDER BY id
        """,
        (venta_id,),
    ).fetchall()
    conn.close()
    return render_template("venta_ticket.html", venta=venta, detalles=detalles)


@app.route("/ventas/cierre")
@login_required
def ventas_cierre():
    fecha = request.args.get("fecha", "").strip() or current_date()
    conn = connect_db()

    resumen = conn.execute(
        """
        SELECT
            COUNT(*) AS ventas,
            COALESCE(SUM(total), 0) AS total,
            COALESCE(SUM(CASE WHEN tipo_pago = 'CONTADO' THEN total ELSE 0 END), 0) AS contado,
            COALESCE(SUM(CASE WHEN tipo_pago = 'FIADO' THEN total ELSE 0 END), 0) AS fiado
        FROM ventas
        WHERE fecha = ?
        """,
        (fecha,),
    ).fetchone()

    pagos = conn.execute(
        """
        SELECT
            COALESCE(SUM(CASE WHEN lower(metodo_pago) = 'efectivo' THEN monto ELSE 0 END), 0) AS efectivo,
            COALESCE(SUM(CASE WHEN lower(metodo_pago) = 'transferencia' THEN monto ELSE 0 END), 0) AS transferencia
        FROM abonos
        WHERE fecha = ?
        """,
        (fecha,),
    ).fetchone()

    por_usuario = conn.execute(
        """
        SELECT
            COALESCE(usuario, '-') AS usuario,
            COUNT(*) AS ventas,
            COALESCE(SUM(total), 0) AS total
        FROM ventas
        WHERE fecha = ?
        GROUP BY usuario
        ORDER BY total DESC
        """,
        (fecha,),
    ).fetchall()

    top_productos = conn.execute(
        """
        SELECT
            vd.sku,
            vd.producto,
            COALESCE(SUM(vd.cantidad), 0) AS unidades,
            COALESCE(SUM(vd.total_linea), 0) AS ingresos
        FROM venta_detalle vd
        JOIN ventas v ON v.id = vd.venta_id
        WHERE v.fecha = ?
        GROUP BY vd.sku, vd.producto
        ORDER BY ingresos DESC
        LIMIT 10
        """,
        (fecha,),
    ).fetchall()
    conn.close()

    return render_template(
        "ventas_cierre.html",
        fecha=fecha,
        resumen=resumen,
        pagos=pagos,
        por_usuario=por_usuario,
        top_productos=top_productos,
    )


@app.route("/admin/backups", methods=["GET", "POST"])
@login_required
@role_required("admin")
def admin_backups():
    ensure_backup_dir()

    if request.method == "POST":
        action = request.form.get("action", "").strip()
        if action == "create":
            target = create_backup()
            flash(f"Respaldo creado: {target.name}", "ok")
        elif action == "restore":
            name = request.form.get("backup_name", "").strip()
            src = BACKUP_DIR / name
            if not src.exists():
                flash("No existe ese respaldo.", "error")
            else:
                shutil.copy2(src, DB_PATH)
                flash(f"Base restaurada desde: {name}", "ok")
        return redirect(url_for("admin_backups"))

    backups = sorted(BACKUP_DIR.glob("inventario_*.db"), reverse=True)
    return render_template("admin_backups.html", backups=backups)


def csv_response(filename: str, headers: list[str], rows: list[list[object]]) -> Response:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(headers)
    writer.writerows(rows)
    return Response(
        out.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.route("/export/inventario.csv")
@login_required
def export_inventario_csv():
    conn = connect_db()
    rows = conn.execute(
        """
        SELECT sku, nombre, categoria, stock_actual, stock_minimo, precio_compra, precio_venta
        FROM productos
        ORDER BY nombre
        """
    ).fetchall()
    conn.close()
    return csv_response(
        "inventario.csv",
        ["SKU", "Producto", "Categoria", "Stock_Actual", "Stock_Minimo", "Precio_Compra", "Precio_Venta"],
        [list(r) for r in rows],
    )


@app.route("/export/ventas.csv")
@login_required
def export_ventas_csv():
    conn = connect_db()
    rows = conn.execute(
        """
        SELECT v.id, v.fecha, v.tipo_pago, COALESCE(c.nombre, '-') AS cliente, v.total, v.usuario
        FROM ventas v
        LEFT JOIN clientes c ON c.id = v.cliente_id
        ORDER BY v.id DESC
        """
    ).fetchall()
    conn.close()
    return csv_response(
        "ventas.csv",
        ["Venta", "Fecha", "Pago", "Cliente", "Total", "Usuario"],
        [list(r) for r in rows],
    )


@app.route("/export/cartera.csv")
@login_required
def export_cartera_csv():
    conn = connect_db()
    rows = conn.execute(
        """
        SELECT
            d.id,
            c.nombre,
            d.fecha,
            d.fecha_vencimiento,
            d.concepto,
            d.monto_total,
            d.monto_pagado,
            (d.monto_total - d.monto_pagado) AS saldo,
            d.estado
        FROM deudas d
        JOIN clientes c ON c.id = d.cliente_id
        ORDER BY d.id DESC
        """
    ).fetchall()
    conn.close()
    return csv_response(
        "cartera.csv",
        ["Deuda", "Cliente", "Fecha", "Vencimiento", "Concepto", "Monto_Total", "Monto_Pagado", "Saldo", "Estado"],
        [list(r) for r in rows],
    )


@app.route("/export/movimientos.csv")
@login_required
def export_movimientos_csv():
    conn = connect_db()
    rows = conn.execute(
        """
        SELECT m.fecha, m.tipo, m.sku, p.nombre, m.cantidad, m.usuario, m.notas
        FROM movimientos m
        JOIN productos p ON p.sku = m.sku
        ORDER BY m.fecha DESC, m.id DESC
        """
    ).fetchall()
    conn.close()
    return csv_response(
        "movimientos.csv",
        ["Fecha", "Tipo", "SKU", "Producto", "Cantidad", "Usuario", "Notas"],
        [list(r) for r in rows],
    )


@app.route("/export/ventas.xlsx")
@login_required
def export_ventas_xlsx():
    try:
        from openpyxl import Workbook
    except Exception:
        flash("Falta dependencia openpyxl. Instala requirements.txt", "error")
        return redirect(url_for("ventas"))

    conn = connect_db()
    rows = conn.execute(
        """
        SELECT v.id, v.fecha, v.tipo_pago, COALESCE(c.nombre, '-') AS cliente, v.total, v.usuario
        FROM ventas v
        LEFT JOIN clientes c ON c.id = v.cliente_id
        ORDER BY v.id DESC
        """
    ).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(["Venta", "Fecha", "Pago", "Cliente", "Total", "Usuario"])
    for row in rows:
        ws.append(list(row))

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return Response(
        stream.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=ventas.xlsx"},
    )


@app.route("/export/cierre.pdf")
@login_required
def export_cierre_pdf():
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
    except Exception:
        flash("Falta dependencia reportlab. Instala requirements.txt", "error")
        return redirect(url_for("ventas_cierre"))

    fecha = request.args.get("fecha", "").strip() or current_date()
    conn = connect_db()
    resumen = conn.execute(
        """
        SELECT
            COUNT(*) AS ventas,
            COALESCE(SUM(total), 0) AS total,
            COALESCE(SUM(CASE WHEN tipo_pago = 'CONTADO' THEN total ELSE 0 END), 0) AS contado,
            COALESCE(SUM(CASE WHEN tipo_pago = 'FIADO' THEN total ELSE 0 END), 0) AS fiado
        FROM ventas WHERE fecha = ?
        """,
        (fecha,),
    ).fetchone()

    pagos = conn.execute(
        """
        SELECT
            COALESCE(SUM(CASE WHEN lower(metodo_pago) = 'efectivo' THEN monto ELSE 0 END), 0) AS efectivo,
            COALESCE(SUM(CASE WHEN lower(metodo_pago) = 'transferencia' THEN monto ELSE 0 END), 0) AS transferencia
        FROM abonos
        WHERE fecha = ?
        """,
        (fecha,),
    ).fetchone()
    conn.close()

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.setTitle(f"Cierre_{fecha}")
    y = 760
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, y, "Cierre Diario de Caja")
    y -= 24
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, y, f"Fecha: {fecha}")
    y -= 24
    pdf.drawString(50, y, f"Ventas: {resumen['ventas']}")
    y -= 18
    pdf.drawString(50, y, f"Total: {format_currency(resumen['total'])}")
    y -= 18
    pdf.drawString(50, y, f"Contado: {format_currency(resumen['contado'])}")
    y -= 18
    pdf.drawString(50, y, f"Fiado: {format_currency(resumen['fiado'])}")
    y -= 18
    pdf.drawString(50, y, f"Efectivo: {format_currency(pagos['efectivo'])}")
    y -= 18
    pdf.drawString(50, y, f"Transferencia: {format_currency(pagos['transferencia'])}")
    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return Response(
        buffer.read(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=cierre_{fecha}.pdf"},
    )


@app.route("/cartera")
@login_required
def cartera():
    conn = connect_db()

    conn.execute(
        """
        UPDATE deudas
        SET estado = 'VENCIDA'
        WHERE estado = 'PENDIENTE'
          AND fecha_vencimiento IS NOT NULL
          AND fecha_vencimiento <> ''
          AND date(fecha_vencimiento) < date('now')
          AND (monto_total - monto_pagado) > 0
        """
    )
    conn.commit()

    clientes_rows = conn.execute(
        """
        SELECT
            c.id,
            c.nombre,
            c.telefono,
            c.correo,
            COALESCE(SUM(d.monto_total - d.monto_pagado), 0) AS saldo_pendiente
        FROM clientes c
        LEFT JOIN deudas d ON d.cliente_id = c.id AND d.estado <> 'PAGADA'
        WHERE c.activo = 1
        GROUP BY c.id, c.nombre, c.telefono, c.correo
        ORDER BY c.nombre
        """
    ).fetchall()

    deudas_rows = conn.execute(
        """
        SELECT
            d.id,
            d.fecha,
            d.fecha_vencimiento,
            d.concepto,
            d.monto_total,
            d.monto_pagado,
            (d.monto_total - d.monto_pagado) AS saldo,
            d.estado,
            c.nombre AS cliente_nombre
        FROM deudas d
        JOIN clientes c ON c.id = d.cliente_id
        ORDER BY CASE WHEN d.estado = 'PENDIENTE' THEN 0 WHEN d.estado = 'VENCIDA' THEN 1 ELSE 2 END,
                 d.fecha DESC,
                 d.id DESC
        """
    ).fetchall()

    producto_concepto_rows = [
        dict(row)
        for row in conn.execute(
            "SELECT sku, nombre, precio_venta FROM productos ORDER BY nombre"
        ).fetchall()
    ]

    total_pendiente = conn.execute(
        """
        SELECT COALESCE(SUM(monto_total - monto_pagado), 0) AS total
        FROM deudas
        WHERE estado <> 'PAGADA'
        """
    ).fetchone()["total"]

    conn.close()
    return render_template(
        "cartera.html",
        clientes=clientes_rows,
        deudas=deudas_rows,
        total_pendiente=total_pendiente,
        productos_concepto=producto_concepto_rows,
    )


@app.route("/cartera/clientes/nuevo", methods=["POST"])
@login_required
@role_required("admin", "supervisor")
def cartera_nuevo_cliente():
    nombre = request.form.get("nombre", "").strip()
    telefono = request.form.get("telefono", "").strip()
    correo = request.form.get("correo", "").strip()
    notas = request.form.get("notas", "").strip()

    if not nombre:
        flash("El nombre del cliente es obligatorio.", "error")
        return redirect(url_for("cartera"))

    conn = connect_db()
    conn.execute(
        "INSERT INTO clientes (nombre, telefono, correo, notas, activo) VALUES (?, ?, ?, ?, 1)",
        (nombre, telefono, correo, notas),
    )
    conn.commit()
    conn.close()
    flash("Cliente agregado a cartera.", "ok")
    return redirect(url_for("cartera"))


@app.route("/cartera/deudas/nueva", methods=["POST"])
@login_required
@role_required("admin", "supervisor")
def cartera_nueva_deuda():
    cliente_id = request.form.get("cliente_id", "").strip()
    concepto = request.form.get("concepto", "").strip()
    fecha_vencimiento = request.form.get("fecha_vencimiento", "").strip()
    monto_total = parse_float(request.form.get("monto_total", "0"), 0.0)

    if not cliente_id:
        flash("Selecciona un cliente.", "error")
        return redirect(url_for("cartera"))
    if not concepto:
        flash("El concepto de la deuda es obligatorio.", "error")
        return redirect(url_for("cartera"))
    if monto_total <= 0:
        flash("El monto total debe ser mayor a 0.", "error")
        return redirect(url_for("cartera"))

    conn = connect_db()
    conn.execute(
        """
        INSERT INTO deudas (
            cliente_id, fecha, fecha_vencimiento, concepto, monto_total,
            monto_pagado, estado, usuario
        ) VALUES (?, ?, ?, ?, ?, 0, 'PENDIENTE', ?)
        """,
        (
            int(cliente_id),
            current_date(),
            fecha_vencimiento or None,
            concepto,
            monto_total,
            session.get("username", ""),
        ),
    )
    conn.commit()
    conn.close()
    flash("Deuda registrada correctamente.", "ok")
    return redirect(url_for("cartera"))


@app.route("/cartera/abonos/nuevo", methods=["POST"])
@login_required
@role_required("admin", "supervisor", "caja")
def cartera_nuevo_abono():
    deuda_id = request.form.get("deuda_id", "").strip()
    metodo_pago = request.form.get("metodo_pago", "").strip()
    monto = parse_float(request.form.get("monto", "0"), 0.0)

    if not deuda_id:
        flash("Selecciona una deuda para abonar.", "error")
        return redirect(url_for("cartera"))
    if monto <= 0:
        flash("El monto de abono debe ser mayor a 0.", "error")
        return redirect(url_for("cartera"))

    conn = connect_db()
    deuda = conn.execute("SELECT * FROM deudas WHERE id = ?", (int(deuda_id),)).fetchone()

    if not deuda:
        conn.close()
        flash("La deuda no existe.", "error")
        return redirect(url_for("cartera"))

    saldo = float(deuda["monto_total"]) - float(deuda["monto_pagado"])
    if saldo <= 0:
        conn.close()
        flash("La deuda ya esta pagada.", "error")
        return redirect(url_for("cartera"))
    if monto > saldo:
        conn.close()
        flash("El abono no puede ser mayor al saldo pendiente.", "error")
        return redirect(url_for("cartera"))

    nuevo_pagado = float(deuda["monto_pagado"]) + monto
    nuevo_saldo = float(deuda["monto_total"]) - nuevo_pagado
    nuevo_estado = "PAGADA" if nuevo_saldo <= 0.00001 else "PENDIENTE"

    conn.execute(
        """
        INSERT INTO abonos (deuda_id, fecha, monto, metodo_pago, usuario)
        VALUES (?, ?, ?, ?, ?)
        """,
        (int(deuda_id), current_date(), monto, metodo_pago, session.get("username", "")),
    )
    conn.execute(
        "UPDATE deudas SET monto_pagado = ?, estado = ? WHERE id = ?",
        (nuevo_pagado, nuevo_estado, int(deuda_id)),
    )
    conn.commit()
    conn.close()
    flash("Abono registrado correctamente.", "ok")
    return redirect(url_for("cartera"))


if __name__ == "__main__":
    bootstrap()
    port = int(os.environ.get("PORT", 5000))
    host = os.environ.get("HOST", "0.0.0.0")
    app.run(host=host, port=port, debug=False)
